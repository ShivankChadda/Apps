#!/usr/bin/env python3
"""Cheap structural checks on a .docx: well-formed XML, python-docx can open it, and
child-element order inside the elements Word is strict about."""
import sys, zipfile
from xml.etree import ElementTree as ET
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
ORDER = {
  'pPr': 'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange'.split(),
  'rPr': 'rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs em lang eastAsianLayout specVanish oMath'.split(),
  'tblPr': 'tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing tblInd tblBorders shd tblLayout tblCellMar tblLook'.split(),
  'tcPr': 'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark'.split(),
  'sectPr': 'headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid'.split(),
  'pBdr': 'top left bottom right between bar'.split(),
  'tcBorders': 'top start left bottom end right insideH insideV tl2br tr2bl'.split(),
  'tblCellMar': 'top start left bottom end right'.split(),
  'lvl': 'start numFmt lvlRestart pStyle isLgl suff lvlText lvlPicBulletId legacy lvlJc pPr rPr'.split(),
}
def check(root, part):
    bad = 0
    for el in root.iter():
        tag = el.tag.replace(W, '')
        if tag in ORDER and el.tag.startswith(W):
            seq = [c.tag.replace(W, '') for c in el if c.tag.startswith(W)]
            idx = []
            for s in seq:
                if s not in ORDER[tag]:
                    print(f'  {part}: <{tag}> has unexpected child <{s}>'); bad += 1; continue
                idx.append(ORDER[tag].index(s))
            if idx != sorted(idx):
                print(f'  {part}: <{tag}> children out of order: {seq}'); bad += 1
    return bad
bad = 0
for f in sys.argv[1:]:
    z = zipfile.ZipFile(f)
    for n in z.namelist():
        if n.endswith('.xml') or n.endswith('.rels'):
            try: root = ET.fromstring(z.read(n))
            except Exception as e: print(f'{f}: {n} is not well-formed: {e}'); bad += 1; continue
            if n.startswith('word/') and n.endswith('.xml'): bad += check(root, n)
    # every paragraph-bearing cell must end with a paragraph
    root = ET.fromstring(z.read('word/document.xml'))
    for tc in root.iter(W + 'tc'):
        kids = [c.tag for c in tc]
        if not kids or kids[-1] != W + 'p': print(f'{f}: table cell does not end with a paragraph'); bad += 1
    try:
        import docx
        d = docx.Document(f)
        print(f'{f}: python-docx ok, {len(d.paragraphs)} paragraphs, {len(d.tables)} tables')
    except Exception as e:
        print(f'{f}: python-docx failed: {e}'); bad += 1
print('PROBLEMS:', bad)
sys.exit(1 if bad else 0)
