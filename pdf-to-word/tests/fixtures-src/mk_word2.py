from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw
import random
random.seed(5)
W = ('the project team reviewed the delivery plan and agreed that milestones should be tied to measurable outcomes rather than calendar dates because dependencies between suppliers change often').split()
def sent(n=14):
    w=[random.choice(W) for _ in range(n)]; w[0]=w[0].capitalize(); return ' '.join(w)+'.'
def para(n=4): return ' '.join(sent(random.randint(9,17)) for _ in range(n))
im = Image.new('RGB', (600, 260), (246, 248, 250)); d = ImageDraw.Draw(im)
for i in range(7): d.rectangle([30+i*76, 220-(30+(i*37)%160), 90+i*76, 220], fill=[(36,85,164),(192,57,43),(46,139,87),(224,161,0)][i%4])
d.line([20,220,580,220], fill=(40,40,40), width=3); im.save('chart2.png')

doc = Document()
s = doc.sections[0]; s.left_margin = s.right_margin = Inches(1.1); s.top_margin = Inches(1); s.bottom_margin = Inches(1)
s.different_first_page_header_footer = True
st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10.5); st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.2
hp = s.header.paragraphs[0]; hp.text = 'Delivery Plan v2 - Confidential'
fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
def field(p, instr):
    r = p.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r._r.append(f1)
    r2 = p.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r2._r.append(it)
    r3 = p.add_run(); f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end'); r3._r.append(f2)
fp.add_run('Page '); field(fp, 'PAGE')
t = doc.add_paragraph('Delivery Plan'); t.style = doc.styles['Title']
doc.add_paragraph('Programme office, second revision').runs[0].italic = True
doc.add_heading('1. Overview', 1)
doc.add_paragraph(para(4)).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
doc.add_heading('1.1 Workstreams', 2)
for txt in ['Procurement', 'Engineering']:
    doc.add_paragraph(txt, style='List Bullet')
    if txt == 'Procurement':
        for sub in ['Framework agreements', 'Spot buys over the threshold']:
            doc.add_paragraph(sub, style='List Bullet 2')
        doc.add_paragraph('Approval by the board', style='List Bullet 3')
    else:
        doc.add_paragraph('Design authority review', style='List Bullet 2')
doc.add_paragraph('Decisions needed:', style='Normal')
for i, txt in enumerate(['Confirm the budget envelope', 'Name the programme sponsor', 'Agree the reporting cadence']):
    doc.add_paragraph(txt, style='List Number')
    if i == 1:
        doc.add_paragraph('Interim sponsor from finance', style='List Number 2')
q = doc.add_paragraph(para(2)); q.style = doc.styles['Quote'] if 'Quote' in [x.name for x in doc.styles] else q.style
code = doc.add_paragraph('def plan(milestones):\n    return sorted(milestones, key=lambda m: m.date)'); 
for r in code.runs: r.font.name = 'Courier New'; r.font.size = Pt(9)
doc.add_heading('2. Schedule', 1)
tb = doc.add_table(rows=4, cols=4); tb.style = 'Table Grid'
tb.cell(0,0).merge(tb.cell(1,0)); tb.cell(0,0).text = 'Phase'
tb.cell(0,1).merge(tb.cell(0,3)); tb.cell(0,1).text = 'Quarter'
for i, h in enumerate(['Q1','Q2','Q3']): tb.cell(1,1+i).text = h
for r, row in enumerate([('Design','done','active','-'),('Build','-','active','active')]):
    for c, v in enumerate(row): tb.cell(2+r, c).text = v
for c in tb.rows[0].cells + tb.rows[1].cells:
    for p in c.paragraphs:
        for run in p.runs: run.bold = True
doc.add_paragraph()
doc.add_picture('chart2.png', width=Inches(4.2)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
cap = doc.add_paragraph('Figure 1. Milestones delivered per quarter'); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap.runs[0].font.size = Pt(8.5)
doc.add_heading('3. Risks', 1)
for _ in range(2): doc.add_paragraph(para(5)).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
ns = doc.add_section(WD_SECTION.NEW_PAGE); ns.orientation = WD_ORIENT.LANDSCAPE; ns.page_width, ns.page_height = ns.page_height, ns.page_width
doc.add_heading('Appendix A. Wide table', 1)
wt = doc.add_table(rows=4, cols=7); wt.style = 'Table Grid'
for c, h in enumerate(['ID','Task','Owner','Start','End','Status','Notes']): wt.cell(0,c).text = h
for r in range(1,4):
    for c, v in enumerate([f'T{r}', 'Install monitoring', 'Ops', '2026-01-1'+str(r), '2026-02-0'+str(r), 'Open', 'Waiting on supplier']): wt.cell(r,c).text = v
doc.save('word2.docx')
